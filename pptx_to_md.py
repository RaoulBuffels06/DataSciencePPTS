from pathlib import Path
from pptx import Presentation

def shape_text(shape) -> str:
    """Extract text from a shape if it has a text frame."""
    if not hasattr(shape, "has_text_frame") or not shape.has_text_frame:
        return ""
    text = []
    for p in shape.text_frame.paragraphs:
        line = p.text.strip()
        if line:
            text.append(line)
    return "\n".join(text).strip()

def pptx_to_markdown(pptx_path: Path, out_dir: Path) -> Path:
    prs = Presentation(str(pptx_path))
    out_dir.mkdir(parents=True, exist_ok=True)

    md_lines = [f"# {pptx_path.stem}", ""]

    for i, slide in enumerate(prs.slides, start=1):
        slide_text_blocks = []

        for shape in slide.shapes:
            t = shape_text(shape)
            if t:
                slide_text_blocks.append(t)

        # Deduplicate while preserving order
        seen = set()
        unique_blocks = []
        for b in slide_text_blocks:
            if b not in seen:
                seen.add(b)
                unique_blocks.append(b)

        md_lines.append(f"## Slide {i}")
        if unique_blocks:
            for block in unique_blocks:
                for line in block.split("\n"):
                    md_lines.append(f"- {line}")
                md_lines.append("")
        else:
            md_lines.append("- (geen tekst op deze slide)")
            md_lines.append("")

    md_path = out_dir / f"{pptx_path.stem}.md"
    md_path.write_text("\n".join(md_lines).strip() + "\n", encoding="utf-8")
    return md_path

def main():
    base = Path(".")
    out_dir = base / "docs" / "slides_as_text"

    pptx_files = sorted(base.glob("*.pptx"))
    if not pptx_files:
        print("Geen .pptx gevonden in deze map.")
        return

    for pptx in pptx_files:
        md = pptx_to_markdown(pptx, out_dir)
        print(f"OK: {pptx.name} -> {md}")

if __name__ == "__main__":
    main()
